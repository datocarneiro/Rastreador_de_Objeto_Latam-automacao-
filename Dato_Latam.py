import io
import pickle
import json
from flask import Flask, render_template, request, session, send_file, Response, stream_with_context
from openpyxl import load_workbook
from werkzeug.utils import secure_filename
import pandas as pd
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
from selenium.common.exceptions import TimeoutException
from webdriver_manager.chrome import ChromeDriverManager

# Configuração do aplicativo Flask
app = Flask(__name__)
app.secret_key = 'dato123'  # Defina uma chave secreta adequada

# # Configuração do Selenium
options = Options()
options.add_argument("--no-sandbox")
options.add_argument("--headless")  # Executar o Chrome de forma oculta
options.add_argument("--disable-dev-shm-usage")

service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)

# Variável global para armazenar a lista de pendentes e o resultado final
lista_pendentes = []
resultado_final_df = None

@app.route('/')
def index():
    return render_template('index.html', pendentes=lista_pendentes)

@app.route('/rastreando')
def rastreando():
    return render_template('rastreando.html', total_pendentes=len(lista_pendentes))

@app.route('/resultado')
def resultado():
    global resultado_final_df
    if resultado_final_df is not None and not resultado_final_df.empty:
        table_html = resultado_final_df.to_html(classes='table table-bordered', index=False)
    else:
        table_html = "<p style='text-align:center;'>Nenhum dado disponível. Por favor, inicie um rastreamento primeiro.</p>"
    return render_template('resultado.html', table_html=table_html)

@app.route('/exportar_excel')
def exportar_excel():
    global resultado_final_df
    if resultado_final_df is None or resultado_final_df.empty:
        return "Nenhum dado disponível para exportação"
    
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
        resultado_final_df.to_excel(writer, index=False, sheet_name='Sheet1')
    output.seek(0)
    
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', download_name='resultado.xlsx', as_attachment=True)

@app.route('/download_modelo')
def download_modelo():
    return send_file('modelo/base_rastreamento.xlsx', as_attachment=True, download_name='base_rastreamento.xlsx')

@app.route('/', methods=['POST'])
def preparar_dados_planilha():
    global lista_pendentes  
    file = request.files['file']

    if not file.filename.endswith('.xlsx'):
        return "Por favor, selecione um arquivo Excel (.xlsx)"

    filename = secure_filename(file.filename)
    file.save(filename)
    session['excel_filename'] = filename  

    planilha = load_workbook(filename)
    aba_ativa = planilha.active

    lista_pendentes = [
        coluna_c.value
        for coluna_a, coluna_c, coluna_d in zip(aba_ativa["A"][1:], aba_ativa["C"][1:], aba_ativa["D"][1:])
        if coluna_d.value != 'ENTREGUE' and coluna_c.value is not None
    ]

    total_pendentes = len(lista_pendentes)
    print("=" * 150)
    print(f'As pendente de entrega são: Total: |{total_pendentes}| {lista_pendentes}')
    print("=" * 150)

    return render_template('index.html', total_pendentes=total_pendentes, pendentes=lista_pendentes)

def captura_status(franquia, awb):
    driver.get(f"https://www.latamcargo.com/en/trackshipment?docNumber={awb}&docPrefix=957&soType=SO")
    wait = WebDriverWait(driver, 30)
    try:
        status_evento = wait.until(EC.visibility_of_element_located((By.XPATH, '//*[@id="statusTable"]/tbody/tr[1]/td[1]')))
        status = status_evento.text
        status_map = {
            'DLV': "Entregue", 'DDL': "Entregue",
            'DDR': "Em rota", 'OFD': "Em rota",
            'DDF': "Fechada",
            'RCF': "Entrada",
            'DEP': "Em transferência", 'ARR': "Em transferência",
            'MAN': "Em transferência", 'BDK': "Em transferência",
            'FOH': "Em transferência", 'RCS': "Em transferência",
            'DDC': "Entrega Cancelada"
        }
        status = f"{status} - {status_map.get(status, 'Desconhecido')}"

        data_evento = wait.until(EC.visibility_of_element_located((By.XPATH, '//*[@id="statusTable"]/tbody/tr[1]/td[6]')))
        data = data_evento.text
        total = len(lista_pendentes)
        captura = f'|TOTAL: {total}|FRANQUIA: {franquia}|AWB: {awb}| STATUS: {status}|DATA_EVENTO: {data}|'
        print("=" * 150)
        print(captura)
        print("=" * 150)

        return status, data
    except TimeoutException:
        print("Erro de tempo limite (latam indisponível)")
        return "Erro de tempo limite", "Erro de tempo limite"

@app.route('/stream')
def stream():
    def generate():
        global resultado_final_df, lista_pendentes
        dados_rastreamento = []
        excel_filename = session.get('excel_filename')

        if excel_filename:
            planilha = load_workbook(excel_filename)
            aba_ativa = planilha.active
            total = len(lista_pendentes)
            current = 0

            for coluna_a, coluna_c, coluna_d in zip(aba_ativa["A"][1:], aba_ativa["C"][1:], aba_ativa["D"][1:]):
                if coluna_d.value != 'ENTREGUE' and coluna_c.value is not None:
                    franquia = coluna_a.value
                    awb = coluna_c.value
                    status, data = captura_status(franquia, awb)
                    dados_rastreamento.append({'FRANQUIA': franquia, 'AWB': awb, 'STATUS': status, 'DATA_EVENTO': data})
                    
                    current += 1
                    # Envia o progresso pro front-end
                    yield f"data: {json.dumps({'current': current, 'total': total, 'awb': awb, 'franquia': franquia})}\n\n"

            resultado_final_df = pd.DataFrame(dados_rastreamento)
            yield f"data: {json.dumps({'done': True})}\n\n"
        else:
            yield f"data: {json.dumps({'error': 'Arquivo não encontrado'})}\n\n"

    return Response(stream_with_context(generate()), mimetype='text/event-stream')

if __name__ == '__main__':
    app.run(host="0.0.0.0", port=8080)
